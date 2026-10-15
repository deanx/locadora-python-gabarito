# relatorio_gerencia.py — a planilha do faturamento, com gráfico, para a gerência

import sys
from pathlib import Path

import matplotlib

matplotlib.use("Agg")  # desenha em arquivo, sem abrir janela

import matplotlib.pyplot as plt
import pandas as pd
from openpyxl import load_workbook
from openpyxl.drawing.image import Image
from openpyxl.styles import Font
from sqlalchemy import create_engine

FORMATO_REAIS = '"R$" #,##0.00'

CONSULTA = """
    SELECT a.inicio, a.valor_final, a.multa, c.nome AS cliente, v.tipo
    FROM alugueis a
    JOIN clientes c ON c.id = a.cliente_id
    JOIN veiculos v ON v.id = a.veiculo_id
"""


def carrega(url_banco: str) -> pd.DataFrame:
    df = pd.read_sql(CONSULTA, create_engine(url_banco), parse_dates=["inicio"])
    df["total"] = df["valor_final"] + df["multa"]
    df["mes"] = df["inicio"].dt.strftime("%Y-%m")
    return df


def desenha_grafico(por_mes: pd.DataFrame, destino: Path) -> None:
    eixo = por_mes.plot(kind="bar", stacked=True, figsize=(9, 4.5), color=["#16324F", "#E0921F"])
    eixo.set_title("Faturamento por mês")
    eixo.set_xlabel("")
    eixo.set_ylabel("R$")
    eixo.legend(title="")
    plt.xticks(rotation=45)
    plt.tight_layout()
    plt.savefig(destino, dpi=100)
    plt.close()


def gera_relatorio(url_banco: str, pasta: Path) -> Path:
    df = carrega(url_banco)
    por_mes = df.pivot_table(index="mes", columns="tipo", values="total", aggfunc="sum").round(2)
    por_mes["Total"] = por_mes.sum(axis=1)
    por_mes.index.name = "Mês"
    clientes = (
        df.groupby("cliente")["total"]
        .agg(["count", "sum"])
        .round(2)
        .rename(columns={"count": "Aluguéis", "sum": "Faturamento"})
        .sort_values("Faturamento", ascending=False)
    )
    clientes.index.name = "Cliente"

    pasta.mkdir(parents=True, exist_ok=True)
    arquivo = pasta / "faturamento.xlsx"
    with pd.ExcelWriter(arquivo) as escritor:
        por_mes.to_excel(escritor, sheet_name="Por mês")
        clientes.to_excel(escritor, sheet_name="Clientes")

    # Acabamento com o openpyxl: reais, larguras, cabeçalho fixo e o gráfico
    planilha = load_workbook(arquivo)
    for aba in planilha.worksheets:
        aba.freeze_panes = "B2"
        aba.column_dimensions["A"].width = 18
        for coluna in "BCDE":
            aba.column_dimensions[coluna].width = 15
        for celula in aba[1]:
            celula.font = Font(bold=True)
    for linha in planilha["Por mês"].iter_rows(min_row=2, min_col=2):
        for celula in linha:
            celula.number_format = FORMATO_REAIS
    for (celula,) in planilha["Clientes"].iter_rows(min_row=2, min_col=3, max_col=3):
        celula.number_format = FORMATO_REAIS

    grafico = pasta / "faturamento.png"
    desenha_grafico(por_mes.drop(columns="Total"), grafico)
    planilha["Por mês"].add_image(Image(str(grafico)), "G2")
    planilha.save(arquivo)
    return arquivo


if __name__ == "__main__":
    url = sys.argv[1] if len(sys.argv) > 1 else "sqlite:///estudo.db"
    print(f"Relatório salvo em {gera_relatorio(url, Path('relatorios'))}")
